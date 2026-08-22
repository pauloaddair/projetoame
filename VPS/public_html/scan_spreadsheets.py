import os
import glob
import csv

dir_path = r"F:\OneDrive\Projeto A.M.E\02_Administrativo_e_Operacional\Planilhas"

print("=== SCANNING CSV FILES IN PLANILHAS DIR ===")
csv_files = glob.glob(os.path.join(dir_path, "*.csv"))

for fpath in csv_files:
    fname = os.path.basename(fpath)
    try:
        with open(fpath, 'r', encoding='utf-8', errors='ignore') as f:
            # Detect delimiter
            sample = f.read(2048)
            f.seek(0)
            delim = ';' if ';' in sample else ','
            
            reader = csv.reader(f, delimiter=delim)
            header = next(reader)
            rows = list(reader)
            
            print(f"\nFile: {fname}")
            print(f" - Rows count: {len(rows)}")
            print(f" - Header columns: {header}")
            
            # Let's inspect date columns
            date_col_idx = -1
            val_col_idx = -1
            for idx, col in enumerate(header):
                col_lower = col.lower()
                if 'data' in col_lower:
                    date_col_idx = idx
                if 'valor' in col_lower or 'valor_realizado' in col_lower:
                    val_col_idx = idx
            
            if date_col_idx != -1 and len(rows) > 0:
                dates = []
                for r in rows:
                    if len(r) > date_col_idx and r[date_col_idx]:
                        dates.append(r[date_col_idx])
                print(f" - Date column '{header[date_col_idx]}' found. Range: {min(dates)} to {max(dates)}")
            if val_col_idx != -1 and len(rows) > 0:
                vals = []
                for r in rows:
                    if len(r) > val_col_idx and r[val_col_idx]:
                        # clean up string
                        val_str = r[val_col_idx].replace('.', '').replace(',', '.')
                        try:
                            vals.append(float(val_str))
                        except:
                            pass
                if vals:
                    print(f" - Sum of values: R$ {sum(vals):,.2f}")
    except Exception as e:
        print(f"Error reading {fname}: {e}")

print("\n=== SCANNING EXCEL FILES IN PLANILHAS DIR ===")
excel_files = glob.glob(os.path.join(dir_path, "*.xlsx"))
for fpath in excel_files:
    fname = os.path.basename(fpath)
    # Since openpyxl might not be installed, let's try to load it
    try:
        import openpyxl
        wb = openpyxl.load_workbook(fpath, read_only=True)
        sheet_names = wb.sheetnames
        print(f"\nFile: {fname}")
        print(f" - Sheets: {sheet_names}")
        for name in sheet_names[:2]: # inspect first two sheets
            sheet = wb[name]
            # estimate rows
            max_r = sheet.max_row
            print(f"   - Sheet '{name}': ~{max_r} rows")
    except Exception as e:
        print(f"\nFile: {fname}")
        print(f" - Can't read Excel (openpyxl might be missing): {e}")
